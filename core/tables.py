"""Čítanie CSV/TSV a zápis TSV pre Anki.

Vstup:  CSV (čiarka alebo stredník zo slovenského Excelu) alebo TSV,
        aj TSV, ktoré vytvorila táto appka (s #hlavičkami).
Výstup: vždy TSV s Anki file headers – import bez mapovania stĺpcov:

    #separator:Tab
    #html:true
    #columns:Front	Back	...	Tags
    #tags column:6

Podporované od Anki 2.1.55.
"""

import csv
import io
import re

CANDIDATE_DELIMITERS = ("\t", ";", ",")
_ANKI_HEADER = re.compile(r"^#[a-z ]+:", re.IGNORECASE)
_SEPARATOR_NAMES = {"tab": "\t", "semicolon": ";", "comma": ",", "space": " ", "pipe": "|", "colon": ":"}


def detect_delimiter(header_line, suffix=""):
    if suffix.lower() == ".tsv":
        return "\t"
    counts = {delimiter: header_line.count(delimiter) for delimiter in CANDIDATE_DELIMITERS}
    best = max(counts, key=counts.get)
    return best if counts[best] > 0 else ","


def _parse_anki_headers(lines):
    """Z riadkov '#kľúč:hodnota' vytiahne separator a columns."""
    headers = {}
    for line in lines:
        key, _, value = line[1:].partition(":")
        headers[key.strip().lower()] = value.rstrip("\r\n")
    return headers


def read_table(path):
    """Vráti (fieldnames, rows, delimiter)."""
    text = path.read_text(encoding="utf-8-sig")
    lines = io.StringIO(text, newline="").readlines()

    header_lines = []
    while lines and _ANKI_HEADER.match(lines[0]):
        header_lines.append(lines.pop(0))
    headers = _parse_anki_headers(header_lines)
    body = "".join(lines)

    if "separator" in headers:
        raw = headers["separator"].strip()
        delimiter = _SEPARATOR_NAMES.get(raw.lower(), raw[:1] or "\t")
    else:
        delimiter = detect_delimiter(lines[0] if lines else "", path.suffix)

    if "columns" in headers:
        fieldnames = next(csv.reader([headers["columns"]], delimiter=delimiter))
        reader = csv.DictReader(io.StringIO(body), fieldnames=fieldnames, delimiter=delimiter)
    else:
        reader = csv.DictReader(io.StringIO(body), delimiter=delimiter)
        fieldnames = reader.fieldnames or []
    rows = list(reader)
    return list(fieldnames), rows, delimiter


def write_anki_tsv(path, fieldnames, rows):
    """TSV s Anki hlavičkami. Stĺpce sa v Anki namapujú podľa mien automaticky."""
    with path.open("w", encoding="utf-8", newline="") as file:
        file.write("#separator:Tab\n")
        file.write("#html:true\n")
        file.write("#columns:" + "\t".join(fieldnames) + "\n")
        if "Tags" in fieldnames:
            file.write(f"#tags column:{fieldnames.index('Tags') + 1}\n")
        writer = csv.DictWriter(
            file, fieldnames=fieldnames, delimiter="\t", extrasaction="ignore", lineterminator="\n"
        )
        writer.writerows(rows)


def output_path(path, suffix_tag):
    """lecture.csv -> lecture_images.tsv (originál sa nikdy neprepíše)."""
    return path.with_name(f"{path.stem}_{suffix_tag}.tsv")


# ------------------------------------------------------------------ XLSX

PREFERRED_SHEETS = ("3 ANKI ALL-LECTURES", "ANKI ALL-LECTURES", "ALL-LECTURES")


def is_xlsx(path):
    return path.suffix.lower() in (".xlsx", ".xlsm")


def xlsx_sheets(path):
    """Názvy listov; odporúčaný list (ALL-LECTURES) je prvý."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        names = list(workbook.sheetnames)
    finally:
        workbook.close()
    preferred = [name for name in names if name.upper() in PREFERRED_SHEETS or "ALL-LECTURES" in name.upper()]
    return preferred + [name for name in names if name not in preferred]


def _cell_text(value):
    """Excel vráti 12, 12.0, '15,16' alebo None -> vždy text ako v CSV."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_xlsx(path, sheet=None):
    """Vráti (fieldnames, rows). Prvý riadok = hlavička, prázdne riadky sa preskočia.

    Číta hodnoty, ktoré Excel uložil pri poslednom uložení (aj výsledky vzorcov).
    """
    from openpyxl import load_workbook

    sheet = sheet or xlsx_sheets(path)[0]
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"List '{sheet}' v {path.name} neexistuje.")
        values = list(workbook[sheet].iter_rows(values_only=True))
    finally:
        workbook.close()

    if not values:
        raise ValueError(f"List '{sheet}' je prázdny.")

    header = [_cell_text(value).strip() for value in values[0]]
    while header and not header[-1]:
        header.pop()
    if not header:
        raise ValueError(f"List '{sheet}' nemá v 1. riadku hlavičku (Front, Back, …).")

    rows = []
    for raw in values[1:]:
        cells = [_cell_text(value) for value in list(raw)[: len(header)]]
        cells += [""] * (len(header) - len(cells))
        if not any(cell.strip() for cell in cells):
            continue
        rows.append(dict(zip(header, cells)))

    if not rows and _has_formulas(path, sheet):
        raise ValueError(
            f"List '{sheet}' má vzorce, ale bez uložených výsledkov.\n"
            "Otvor súbor v Exceli, ulož ho (Ctrl+S) a skús znova."
        )
    return header, rows


def _has_formulas(path, sheet):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True)
    try:
        for row in workbook[sheet].iter_rows(min_row=2, max_row=20, values_only=True):
            if any(isinstance(value, str) and value.startswith("=") for value in row):
                return True
    finally:
        workbook.close()
    return False


def read_any(path, sheet=None):
    """CSV, TSV aj XLSX -> (fieldnames, rows)."""
    if is_xlsx(path):
        return read_xlsx(path, sheet)
    fieldnames, rows, _delimiter = read_table(path)
    return fieldnames, rows
