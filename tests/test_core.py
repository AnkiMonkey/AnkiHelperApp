import pytest

from core import fields, media, naming, tables
from core.pages import parse_cell_pages, parse_pages


# ------------------------------------------------------------ pages

@pytest.mark.parametrize(
    "value, expected",
    [
        ("12", [12]),
        ("12.0", [12]),
        (" 15 ", [15]),
        ("1,2,10,99", [1, 2, 10, 99]),
        ("1;2", [1, 2]),
        ("1 2", [1, 2]),
        ("5-7", [5, 6, 7]),
        ("1, 5 – 7", [1, 5, 6, 7]),
        ("7-5", [5, 6, 7]),
        ("3,3,4", [3, 4]),
    ],
)
def test_cell_pages_valid(value, expected):
    assert parse_cell_pages(value) == expected


@pytest.mark.parametrize(
    "value",
    [None, "", "   ", "abc", "12 strana", '<img src="x.jpg">', "0", "1-9999", "12.5", "-3"],
)
def test_cell_pages_invalid(value):
    assert parse_cell_pages(value) is None


def test_parse_pages_for_pdf():
    assert parse_pages("1,3,5-7", 10) == {0, 2, 4, 5, 6}
    with pytest.raises(ValueError):
        parse_pages("11", 10)
    with pytest.raises(ValueError):
        parse_pages("a", 10)


# ----------------------------------------------------------- naming

def test_naming_matches_export_and_tags():
    base = naming.lecture_base(" PharmII ", "P", 1)
    assert base == "PharmII_P_01"
    assert naming.page_filename(base, 3) == "PharmII_P_01_S_03.jpg"
    assert naming.page_filename(base, 120) == "PharmII_P_01_S_120.jpg"


def test_image_tag_exact_format():
    assert naming.image_tag("anat_P_08_S_09.jpg") == '<img src="anat_P_08_S_09.jpg">'
    assert naming.images_html("anat_P_08", [9, 10]) == '<img src="anat_P_08_S_09.jpg"><br><img src="anat_P_08_S_10.jpg">'


def test_duplicate_bases():
    assert naming.find_duplicate_bases(["a", "b", "a", "a"]) == ["a"]
    assert naming.find_duplicate_bases(["a", "b"]) == []


def test_validate_base():
    assert naming.validate_base(" Memorix ") == "Memorix"
    with pytest.raises(ValueError):
        naming.validate_base("bad/name")
    with pytest.raises(ValueError):
        naming.validate_base("  ")


# ----------------------------------------------------------- fields

def test_fix_back_field_and_idempotent():
    raw = "O: crista iliaca U: trochanter major I: n. gluteus F: abdukcia"
    fixed = fields.fix_back_field(raw)
    assert fixed == (
        "<b>O:</b> crista iliaca<br><b>U:</b> trochanter major"
        "<br><b>I:</b> n. gluteus<br><b>F:</b> abdukcia"
    )
    assert fields.fix_back_field(fixed) == fixed


def test_add_images_multi_page_and_skips():
    rows = [
        {"Front": "a", "Source": "12", "Personal Notes": ""},
        {"Front": "b", "Source": "15,16", "Personal Notes": "pozri skriptá"},
        {"Front": "c", "Source": '<img src="old.jpg">', "Personal Notes": "1-2"},
    ]
    changed, skipped = fields.add_images_to_rows(rows, ["Source", "Personal Notes"], "PharmII_P_01")
    assert changed == 3
    assert rows[0]["Source"].count("<img") == 1
    assert rows[1]["Source"].count("<img") == 2 and "<br>" in rows[1]["Source"]
    assert "PharmII_P_01_S_16.jpg" in rows[1]["Source"]
    assert rows[1]["Personal Notes"] == "pozri skriptá"  # text ostal
    assert rows[2]["Source"] == '<img src="old.jpg">'  # hotový tag nehlási ako chybu
    assert skipped == [(3, "Personal Notes", "pozri skriptá")]


def test_fill_empty_tags():
    rows = [{"Tags": ""}, {"Tags": "x"}, {"Tags": None}]
    assert fields.fill_empty_tags(rows, "Pharmazie::PharmII::V01") == 2
    assert rows[1]["Tags"] == "x"


# ----------------------------------------------------------- tables

@pytest.mark.parametrize("delimiter, suffix", [(",", ".csv"), (";", ".csv"), ("\t", ".tsv"), ("\t", ".csv")])
def test_read_any_input_write_anki_tsv(tmp_path, delimiter, suffix):
    path = tmp_path / f"deck{suffix}"
    header = delimiter.join(["Front", "Back", "Source", "Tags"])
    line = delimiter.join(["MU heparín", "viaže sa na antitrombín", "15", "Pharmazie::PharmII::V01"])
    path.write_text(f"\ufeff{header}\n{line}\n", encoding="utf-8")

    fieldnames, rows, detected = tables.read_table(path)
    assert detected == delimiter
    assert fieldnames == ["Front", "Back", "Source", "Tags"]
    assert rows[0]["Front"] == "MU heparín"

    out = tables.output_path(path, "images")
    assert out.name == "deck_images.tsv"
    rows[0]["Source"] = '<img src="PharmII_P_01_S_15.jpg">'
    tables.write_anki_tsv(out, fieldnames, rows)

    text = out.read_text(encoding="utf-8")
    assert text.startswith("#separator:Tab\n#html:true\n#columns:Front\tBack\tSource\tTags\n#tags column:4\n")
    # výstup appky sa dá znova načítať (reťazenie: images -> tagged -> fixed)
    again_fields, again_rows, again_delimiter = tables.read_table(out)
    assert again_delimiter == "\t"
    assert again_fields == fieldnames
    assert again_rows == rows


def test_tsv_multiline_and_special_fields(tmp_path):
    out = tmp_path / "x.tsv"
    rows = [{"Front": "a\tb", "Back": "riadok 1\nriadok 2", "Tags": ""}]
    tables.write_anki_tsv(out, ["Front", "Back", "Tags"], rows)
    assert tables.read_table(out)[1] == rows


def test_csv_header_starting_with_hash_is_not_anki_header(tmp_path):
    path = tmp_path / "index.csv"
    path.write_text("#;topic\n1;krv\n", encoding="utf-8")
    fieldnames, rows, _ = tables.read_table(path)
    assert fieldnames == ["#", "topic"] and rows[0]["topic"] == "krv"


def test_table_with_comma_inside_quoted_field(tmp_path):
    path = tmp_path / "deck.csv"
    path.write_text('Front;Back\n"a";"x, y, z"\n', encoding="utf-8")
    _, rows, delimiter = tables.read_table(path)
    assert delimiter == ";"
    assert rows[0]["Back"] == "x, y, z"


# ------------------------------------------------------------ media

def test_media_plan_and_transfer(tmp_path):
    src, dst = tmp_path / "src", tmp_path / "collection.media"
    src.mkdir()
    dst.mkdir()
    (src / "new.jpg").write_bytes(b"new")
    (src / "same.jpg").write_bytes(b"same")
    (src / "changed.jpg").write_bytes(b"v2")
    (src / "notes.txt").write_text("not an image")
    (dst / "same.jpg").write_bytes(b"same")
    (dst / "changed.jpg").write_bytes(b"v1")

    files = media.image_files(src)
    assert [f.name for f in files] == ["changed.jpg", "new.jpg", "same.jpg"]

    plan = media.plan_transfer(files, dst)
    assert [f.name for f in plan.new] == ["new.jpg"]
    assert [f.name for f in plan.identical] == ["same.jpg"]
    assert [f.name for f in plan.changed] == ["changed.jpg"]
    assert "changed.jpg" in media.describe_plan(plan)

    # bez prepisu
    assert media.execute_transfer(plan, overwrite_changed=False) == 1
    assert (dst / "changed.jpg").read_bytes() == b"v1"

    # s prepisom (nový plán)
    plan = media.plan_transfer(files, dst)
    assert media.execute_transfer(plan, overwrite_changed=True) == 1
    assert (dst / "changed.jpg").read_bytes() == b"v2"


def test_media_move(tmp_path):
    src, dst = tmp_path / "src", tmp_path / "dst"
    src.mkdir()
    (src / "a.jpg").write_bytes(b"a")
    plan = media.plan_transfer(media.image_files(src), dst)
    assert media.execute_transfer(plan, move=True) == 1
    assert not (src / "a.jpg").exists() and (dst / "a.jpg").exists()


# ------------------------------------------------------------- xlsx

def _make_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "1 PREHĽAD"
    ws.append(["#", "topic"])
    lec = wb.create_sheet("V01")
    lec.append(["Front", "Back", "Personal Notes", "Source", "Tags", None])
    lec.append(["MU heparín", "AT III", 15, 15.0, "Pharmazie::PharmII::V01"])
    lec.append([None, None, None, None, None])  # prázdny riadok sa preskočí
    lec.append(["CESTA POD heparín", "i.v.", None, "15,16", None])
    lec.append(["NÚ", "HIT", "pozri skriptá", 16, None])
    allx = wb.create_sheet("3 ANKI ALL-LECTURES")
    allx.append(["Front", "Back"])
    allx.append(["a", "b"])
    wb.save(path)


def test_xlsx_sheets_preferred_first(tmp_path):
    path = tmp_path / "PharmII.xlsx"
    _make_xlsx(path)
    assert tables.xlsx_sheets(path)[0] == "3 ANKI ALL-LECTURES"
    fieldnames, rows = tables.read_any(path)  # bez listu -> odporúčaný
    assert fieldnames == ["Front", "Back"] and rows == [{"Front": "a", "Back": "b"}]


def test_xlsx_read_lecture_sheet_and_images(tmp_path):
    path = tmp_path / "PharmII.xlsx"
    _make_xlsx(path)
    fieldnames, rows = tables.read_any(path, "V01")
    assert fieldnames == ["Front", "Back", "Personal Notes", "Source", "Tags"]
    assert len(rows) == 3
    assert rows[0]["Personal Notes"] == "15" and rows[0]["Source"] == "15"  # nie 15.0
    assert rows[1]["Source"] == "15,16" and rows[1]["Tags"] == ""

    changed, skipped = fields.add_images_to_rows(rows, ["Personal Notes", "Source"], "anat_P_08")
    assert rows[0]["Source"] == '<img src="anat_P_08_S_15.jpg">'
    assert rows[1]["Source"] == '<img src="anat_P_08_S_15.jpg"><br><img src="anat_P_08_S_16.jpg">'
    assert changed == 4 and skipped == [(4, "Personal Notes", "pozri skriptá")]

    out = tables.output_path(path, "images")
    assert out.name == "PharmII_images.tsv"
    tables.write_anki_tsv(out, fieldnames, rows)
    assert tables.read_table(out)[1] == rows


def test_xlsx_in_table_list_without_lock_files(tmp_path):
    from core.system import list_tables
    _make_xlsx(tmp_path / "a.xlsx")
    (tmp_path / "~$a.xlsx").write_bytes(b"lock")
    (tmp_path / "b.csv").write_text("Front\nx\n")
    assert [p.name for p in list_tables(tmp_path)] == ["a.xlsx", "b.csv"]


def test_xlsx_formulas_without_cached_values(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "f.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Front", "Back"])
    ws.append(["=A1", "=B1"])
    wb.save(path)
    with pytest.raises(ValueError, match="Ctrl\\+S"):
        tables.read_any(path)
